#!/usr/bin/env python3

# this scipt is entirely ai generated

import argparse
import os
import re
import unicodedata
from pathlib import Path
from typing import Optional, Tuple

try:
    import psycopg2
except ModuleNotFoundError as exc:
    raise SystemExit(
        "Missing dependency: psycopg2-binary. Install it with: pip3 install psycopg2-binary"
    ) from exc

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install it with: pip3 install openpyxl"
    ) from exc


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
MAX_EMAIL_LEN = 255
MAX_NAME_LEN = 255
MAX_TVA_LEN = 255
MAX_STATUS_LEN = 16


def normalize_email(value: object) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    match = EMAIL_RE.search(text)
    if not match:
        return None
    return match.group(0).lower()


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text


def first_header_index(header_map: dict[str, int], *names: str) -> Optional[int]:
    for name in names:
        if name in header_map:
            return header_map[name]
    return None


def clean_text(value: object) -> str:
    text = str(value or "").strip()
    if len(text) >= 2 and text[0] == text[-1] and text[0] in ("'", '"'):
        text = text[1:-1].strip()
    return " ".join(text.split())


def split_name(display_name: str) -> Tuple[str, str]:
    cleaned = clean_text(display_name)
    if not cleaned:
        return "", ""
    parts = cleaned.split(" ")
    return parts[0], " ".join(parts[1:])


def truncate(value: str, max_len: int) -> str:
    return value[:max_len] if len(value) > max_len else value


def build_truncation_warnings(
    contacts: list[tuple[str, str, str]]
) -> list[str]:
    warnings: list[str] = []
    for email, display_name, vat_number in contacts:
        first_name, last_name = split_name(display_name)
        checks = [
            ("email", email, MAX_EMAIL_LEN),
            ("first_name", first_name, MAX_NAME_LEN),
            ("last_name", last_name, MAX_NAME_LEN),
            ("tva", vat_number, MAX_TVA_LEN),
        ]
        for field_name, value, max_len in checks:
            truncated = truncate(value, max_len)
            if value != truncated:
                warnings.append(
                    f"[WARN] {email} {field_name} will be truncated "
                    f"({len(value)} -> {max_len}): '{value}' -> '{truncated}'"
                )
    return warnings


def list_contact_rows(xlsx_path: Path) -> list[tuple[str, str, str]]:
    wb = load_workbook(filename=xlsx_path, data_only=True, read_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)

    if not header:
        wb.close()
        return []

    header_map: dict[str, int] = {}
    for idx, col_name in enumerate(header):
        normalized = normalize_header(col_name)
        if normalized:
            header_map[normalized] = idx

    email_col_idx = first_header_index(header_map, "e-mail", "email")
    display_name_col_idx = first_header_index(header_map, "nom d'affichage")
    vat_col_idx = first_header_index(header_map, "n° tva", "n tva")

    if email_col_idx is None:
        wb.close()
        raise SystemExit("Column not found: E-mail")

    seen: set[tuple[str, str, str]] = set()
    contacts: list[tuple[str, str, str]] = []

    for row in rows:
        if not row:
            continue

        email_value = row[email_col_idx] if email_col_idx < len(row) else None
        email = normalize_email(email_value)
        if not email:
            continue

        display_name = ""
        if display_name_col_idx is not None and display_name_col_idx < len(row):
            display_name = clean_text(row[display_name_col_idx])

        vat_number = ""
        if vat_col_idx is not None and vat_col_idx < len(row):
            vat_number = clean_text(row[vat_col_idx])

        contact = (email, display_name, vat_number)
        if contact not in seen:
            seen.add(contact)
            contacts.append(contact)

    wb.close()
    return contacts


def load_dotenv_if_present() -> None:
    root_env = Path(__file__).resolve().parents[2] / ".env"
    if not root_env.exists():
        return

    for line in root_env.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def db_connection_candidates() -> list[tuple[str, str]]:
    required = [
        "POSTGRES_USER",
        "POSTGRES_PASSWORD",
        "POSTGRES_HOST",
        "POSTGRES_PORT",
        "POSTGRES_DB",
    ]
    missing = [key for key in required if not os.getenv(key)]
    if missing:
        raise SystemExit(f"Missing DB env vars: {', '.join(missing)}")

    user = os.environ["POSTGRES_USER"]
    password = os.environ["POSTGRES_PASSWORD"]
    host = os.environ["POSTGRES_HOST"]
    port = os.environ["POSTGRES_PORT"]
    db_name = os.environ["POSTGRES_DB"]
    hosts = [host]
    if host == "postgres":
        hosts.extend(["localhost", "127.0.0.1"])

    candidates: list[tuple[str, str]] = []
    for candidate_host in hosts:
        dsn = (
            f"postgresql://{user}:{password}@{candidate_host}:{port}/{db_name}"
            "?sslmode=disable"
        )
        candidates.append((candidate_host, dsn))
    return candidates


def upsert_users(contacts: list[tuple[str, str, str]]) -> int:
    query = """
        INSERT INTO users (email, password, first_name, last_name, tva, status)
        VALUES (%s, %s, %s, %s, %s, %s)
        ON CONFLICT (email) DO UPDATE SET
            password = EXCLUDED.password,
            first_name = EXCLUDED.first_name,
            last_name = EXCLUDED.last_name,
            tva = EXCLUDED.tva,
            status = EXCLUDED.status
    """
    last_error: Optional[Exception] = None
    for host, dsn in db_connection_candidates():
        try:
            with psycopg2.connect(dsn) as conn:
                with conn.cursor() as cur:
                    for email, display_name, vat_number in contacts:
                        first_name, last_name = split_name(display_name)
                        email_db = truncate(email, MAX_EMAIL_LEN)
                        first_name_db = truncate(first_name, MAX_NAME_LEN)
                        last_name_db = truncate(last_name, MAX_NAME_LEN)
                        vat_db = truncate(vat_number, MAX_TVA_LEN)
                        status_db = truncate("active", MAX_STATUS_LEN)
                        cur.execute(
                            query,
                            (
                                email_db,
                                "",
                                first_name_db,
                                last_name_db,
                                vat_db,
                                status_db,
                            ),
                        )
            return len(contacts)
        except psycopg2.OperationalError as exc:
            last_error = exc
            continue
    raise SystemExit(f"Database connection failed: {last_error}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Import contacts from Contact.xlsx to users table."
    )
    parser.add_argument(
        "file",
        nargs="?",
        default="Contact.xlsx",
        help="Path to Contact.xlsx (default: Contact.xlsx)",
    )
    parser.add_argument(
        "--print-only",
        action="store_true",
        help="Print parsed values without writing to database.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Maximum number of contacts to process (0 means no limit).",
    )
    args = parser.parse_args()

    xlsx_path = Path(args.file)
    if not xlsx_path.exists():
        raise SystemExit(f"File not found: {xlsx_path}")

    contacts = list_contact_rows(xlsx_path)
    if not contacts:
        print("No email found.")
        return

    if args.limit < 0:
        raise SystemExit("--limit must be >= 0")
    if args.limit > 0:
        contacts = contacts[: args.limit]

    load_dotenv_if_present()

    for email, display_name, vat_number in contacts:
        print(f"{email} | {display_name} | {vat_number}")

    truncation_warnings = build_truncation_warnings(contacts)
    if truncation_warnings:
        for warning in truncation_warnings:
            print(warning)
        print(f"[WARN] {len(truncation_warnings)} field value(s) will be truncated.")

    if args.print_only:
        return

    imported_count = upsert_users(contacts)
    print(f"Imported {imported_count} users into database.")


if __name__ == "__main__":
    main()
